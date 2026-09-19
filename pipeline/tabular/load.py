"""The generic CSV/spreadsheet loader — ADDING_DATA.md's "tractable piece", built.

Every prior source got its own module because every source joins to the spine by a
different key. This one is generic because the *human* supplies what the bespoke modules
hardcoded: which column is the venue, which is the operator, which is the date, and a
dataset-level source citation. No LLM touches this path and nothing here costs money —
mapping is dropdowns, joining is the same `find_candidates` the article path uses.

The join keeps the rule the labor loader taught: **a name match alone never joins a row
to the spine.** Every join needs a second corroborating gate — the row's state matches
the candidate's state, or the row's year falls inside the candidate's operating window.
A mapping with neither a state nor a date column cannot join anything, and says so
instead of joining on names anyway.

Every exclusion is tagged with a reason and kept. Rows are never dropped: they come out
as `matched`, `ambiguous`, `no_spine_match`, `not_corroborated`, or `rejected` (failed
schema validation), each browsable with its raw cells.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
from pathlib import Path
from typing import Any

from ..extract.candidates import build_index, find_candidates
from ..extract.run import mint_event_id
from ..schema import EVENT_TYPES, normalize_operator, validate_event
from ..spans.gold import open_at

SAMPLE_ROWS = 8
MAX_ROWS = 5000  # a spreadsheet bigger than this deserves its own module and a conversation

US_STATES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR", "california": "CA",
    "colorado": "CO", "connecticut": "CT", "delaware": "DE", "florida": "FL", "georgia": "GA",
    "hawaii": "HI", "idaho": "ID", "illinois": "IL", "indiana": "IN", "iowa": "IA",
    "kansas": "KS", "kentucky": "KY", "louisiana": "LA", "maine": "ME", "maryland": "MD",
    "massachusetts": "MA", "michigan": "MI", "minnesota": "MN", "mississippi": "MS",
    "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV",
    "new hampshire": "NH", "new jersey": "NJ", "new mexico": "NM", "new york": "NY",
    "north carolina": "NC", "north dakota": "ND", "ohio": "OH", "oklahoma": "OK",
    "oregon": "OR", "pennsylvania": "PA", "rhode island": "RI", "south carolina": "SC",
    "south dakota": "SD", "tennessee": "TN", "texas": "TX", "utah": "UT", "vermont": "VT",
    "virginia": "VA", "washington": "WA", "west virginia": "WV", "wisconsin": "WI",
    "wyoming": "WY", "district of columbia": "DC",
}
_ABBREVS = set(US_STATES.values())


def read_table(path: Path) -> tuple[list[str], list[list[str]]]:
    """Headers + rows out of a .csv or .xlsx. Cells come back as stripped strings."""
    if path.suffix.lower() == ".csv":
        # utf-8-sig first: Excel writes a BOM, and a BOM inside the first header name
        # makes that column unmappable while looking identical on screen.
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                with path.open(newline="", encoding=enc) as fh:
                    rows = [[(c or "").strip() for c in r] for r in csv.reader(fh)]
                break
            except UnicodeDecodeError:
                continue
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(["" if c is None else str(c).strip() for c in r])
        wb.close()
    else:
        raise SystemExit(f"not a table format: {path.suffix}")

    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        raise SystemExit(f"{path.name}: no rows")
    if len(rows) - 1 > MAX_ROWS:
        raise SystemExit(f"{path.name}: {len(rows) - 1} rows is past the {MAX_ROWS}-row "
                         f"limit for the web path — this dataset deserves its own module")
    headers = rows[0]
    return headers, rows[1:]


def inspect(path: Path) -> dict[str, Any]:
    headers, rows = read_table(path)
    return {"file": path.name, "headers": headers,
            "sample_rows": rows[:SAMPLE_ROWS], "total_rows": len(rows)}


def parse_cell_date(raw: str) -> tuple[str | None, str | None]:
    """(YYYY-MM-DD, precision) from a spreadsheet cell, or (None, None).

    Year-only cells become Jan 1 with precision=year, which is the workbook convention
    `schema._check_precision` enforces. Nothing here guesses: an unreadable cell is
    (None, None), which downstream turns into needs_review, not into a date.
    """
    s = (raw or "").strip()
    if not s:
        return None, None
    if re.fullmatch(r"\d{4}", s):
        return f"{s}-01-01", "year"
    if re.fullmatch(r"\d{4}-\d{2}", s):
        return f"{s}-01", "month"
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(?:[T ].*)?", s)
    if m:
        try:
            return dt.date(int(m[1]), int(m[2]), int(m[3])).isoformat(), "exact"
        except ValueError:
            return None, None
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        try:
            return dt.date(int(m[3]), int(m[1]), int(m[2])).isoformat(), "exact"
        except ValueError:
            return None, None
    from ..ingest.parse import parse_date  # "March 3, 2005" and friends
    iso = parse_date(s)
    return (iso, "exact") if iso else (None, None)


def norm_state(raw: str) -> str | None:
    s = (raw or "").strip()
    if not s:
        return None
    if s.upper() in _ABBREVS:
        return s.upper()
    return US_STATES.get(s.lower())


def _cell(row: list[str], idx: int | None) -> str:
    if idx is None or idx < 0 or idx >= len(row):
        return ""
    return row[idx]


def map_rows(path: Path, mapping: dict[str, Any], *, batch_id: str,
             index: dict[str, Any] | None = None) -> dict[str, Any]:
    """Run the mapping over every row. Deterministic, free, re-runnable.

    `mapping` (column values are header names, or absent):
      venue_col (required), operator_col (required), date_col, event_type_col, state_col,
      author_col, url_col,
      default_event_type, source_citation (required), source_date (required, YYYY-MM-DD),
      dataset_title, source_url

    Citation capture: `author_col` / `url_col` cells are recorded exactly as typed into
    `extras.citation`; `source_url` is the dataset-level fallback when no per-row URL
    column exists. The access date is stamped automatically. Absent stays absent.
    """
    headers, rows = read_table(path)
    col = {h: i for i, h in enumerate(headers)}

    def idx(key: str) -> int | None:
        name = mapping.get(key)
        if not name:
            return None
        if name not in col:
            raise SystemExit(f"mapping names column {name!r} which is not in {path.name}")
        return col[name]

    venue_i, operator_i = idx("venue_col"), idx("operator_col")
    date_i, type_i, state_i = idx("date_col"), idx("event_type_col"), idx("state_col")
    author_i, url_i = idx("author_col"), idx("url_col")
    dataset_url = (mapping.get("source_url") or "").strip()
    accessed = dt.datetime.now(dt.timezone.utc).date().isoformat()
    if venue_i is None or operator_i is None:
        raise SystemExit("mapping needs venue_col and operator_col")

    citation = (mapping.get("source_citation") or "").strip()
    source_date = (mapping.get("source_date") or "").strip()
    if not citation:
        raise SystemExit("a dataset-level source citation is required — rows with no "
                         "provenance cannot become events")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", source_date):
        raise SystemExit(f"source_date {source_date!r} must be YYYY-MM-DD")
    title = (mapping.get("dataset_title") or "").strip() or path.name

    default_type = (mapping.get("default_event_type") or "").strip() or None
    if type_i is None and default_type is None:
        raise SystemExit("map an event-type column or pick a default event type")
    if default_type and default_type not in EVENT_TYPES:
        raise SystemExit(f"default_event_type {default_type!r} not in {EVENT_TYPES}")

    # The corroboration rule, stated up front rather than discovered per row: with neither
    # a state column nor a date column mapped, no row can pass the second gate, so refuse
    # the whole mapping instead of emitting 500 identical exclusions.
    if state_i is None and date_i is None:
        raise SystemExit("a name match alone never joins a row to the spine — map a state "
                         "column or a date column so every join has a second gate")

    if index is None:
        index = build_index()

    rel_file = str(path.resolve().relative_to(Path(__file__).resolve().parent.parent))
    events: list[dict[str, Any]] = []
    results: list[dict[str, Any]] = []
    counts = {"matched": 0, "ambiguous": 0, "no_spine_match": 0,
              "not_corroborated": 0, "rejected": 0}

    for n, row in enumerate(rows, 2):  # 2 = first data row in the spreadsheet's own numbering
        venue_text = _cell(row, venue_i)
        operator_raw = _cell(row, operator_i)
        res: dict[str, Any] = {"row": n, "venue": venue_text, "operator": operator_raw,
                               "cells": dict(zip(headers, row))}

        if not venue_text or not operator_raw:
            res.update(outcome="rejected",
                       reason="venue or operator cell is blank")
            counts["rejected"] += 1
            results.append(res)
            continue

        event_date, precision = (None, None)
        if date_i is not None:
            event_date, precision = parse_cell_date(_cell(row, date_i))
        year = int(event_date[:4]) if event_date else None
        row_state = norm_state(_cell(row, state_i)) if state_i is not None else None

        etype = default_type
        if type_i is not None:
            raw_type = _cell(row, type_i).strip().lower()
            if raw_type in EVENT_TYPES:
                etype = raw_type
            elif raw_type:
                res.update(outcome="rejected",
                           reason=f"event type {raw_type!r} not in {EVENT_TYPES}")
                counts["rejected"] += 1
                results.append(res)
                continue
        if not etype:
            res.update(outcome="rejected", reason="no event type and no default")
            counts["rejected"] += 1
            results.append(res)
            continue

        candidates = find_candidates(venue_text, index, year)
        if not candidates:
            res.update(outcome="no_spine_match",
                       reason="no spine venue matches this name — real row, off this map")
            counts["no_spine_match"] += 1
            results.append(res)
            continue

        # Gate 2: state or year must also match. Name alone never joins. The year gate
        # requires the venue to actually carry a dated operating window — `open_at` counts
        # unknown dates as plausible, and "the year does not contradict a window nobody
        # knows" is a name-only match wearing a second gate's costume.
        corroborated = []
        for c in candidates:
            gates = []
            if row_state and c.get("state") == row_state:
                gates.append("state")
            if (year is not None and c.get("plausible_at_article_date")
                    and (c.get("opened_date") or c.get("closed_date"))):
                gates.append("year")
            if gates:
                corroborated.append((c, gates))

        if not corroborated:
            have = "state" if row_state else ("year" if year else "nothing usable")
            res.update(outcome="not_corroborated",
                       reason=f"name matched {len(candidates)} venue(s) but the second gate "
                              f"failed (row offers {have}); name alone never joins",
                       candidates=[c["canonical_name"] for c in candidates[:5]])
            counts["not_corroborated"] += 1
            results.append(res)
            continue
        if len(corroborated) > 1:
            res.update(outcome="ambiguous",
                       reason=f"{len(corroborated)} venues corroborate — a person must pick",
                       candidates=[c["canonical_name"] for c, _ in corroborated[:5]])
            counts["ambiguous"] += 1
            results.append(res)
            continue

        cand, gates = corroborated[0]
        norm = normalize_operator(operator_raw)
        ev = {
            "operator": operator_raw,
            "operator_normalized": norm["operator_normalized"],
            "sub_brand": norm.get("sub_brand"),
            "venue_id": cand["venue_id"],
            "institution": None,
            "venue_name_as_written": venue_text,
            "event_type": etype,
            "event_date": event_date,
            "date_precision": precision or "approx",
            "first_outsourcing": None,
            "contract_value_usd": None,
            "contract_length_years": None,
            "losing_bidders": None,
            "source_publication": citation,
            "source_date": source_date,
            "source_title": title,
            "source_file": rel_file,
            "extraction_confidence": 1.0,
            "needs_review": bool(norm.get("needs_review")) or event_date is None,
            "notes": f"csv row {n}; joined by name+{'+'.join(gates)}",
            "extras": {"batch": batch_id, "row": n, "join_gates": gates},
        }
        # Citation: the row's own cells, exactly as typed, else the dataset-level URL.
        # Found keys only — a blank cell writes nothing, and nothing is derived.
        cit: dict[str, str] = {}
        row_author = _cell(row, author_i).strip() if author_i is not None else ""
        row_url = _cell(row, url_i).strip() if url_i is not None else ""
        if row_author:
            cit["author"] = row_author
        if row_url or dataset_url:
            cit["url"] = row_url or dataset_url
        cit["accessed"] = accessed
        ev["extras"]["citation"] = cit
        ev["event_id"] = mint_event_id(
            {"source_file": rel_file, "source_offset": n}, ev)
        problems = validate_event(ev)
        if problems:
            res.update(outcome="rejected", reason="; ".join(problems))
            counts["rejected"] += 1
            results.append(res)
            continue

        res.update(outcome="matched", venue_matched=cand["canonical_name"],
                   venue_id=cand["venue_id"], gates=gates, event_id=ev["event_id"])
        counts["matched"] += 1
        results.append(res)
        events.append(ev)

    return {"file": path.name, "mapping": mapping, "counts": counts,
            "total_rows": len(rows), "results": results, "events": events}
